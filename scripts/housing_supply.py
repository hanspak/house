"""국토부 공식 미분양 엑셀에서 전체·준공 후 미분양을 함께 추출한다."""
import datetime as dt
import gzip
import json
from pathlib import Path
import re
import urllib.parse
import urllib.request
import openpyxl

from trade_regions import PROVINCES, districts

ROOT = Path(__file__).resolve().parents[1]
SOURCE = 'https://stat.molit.go.kr/portal/cate/statMetaView.do?hRsId=32'


def norm(value):
    return re.sub(r'\s+', '', str(value or '')).replace('(구)', '')


def region_key(province, district):
    province, district = norm(province), norm(district)
    if province == '전남광주':
        if district == '계':
            return None
        province = '광주' if district in ('동구', '서구', '남구', '북구', '광산구') else '전남'
    if province == '전국':
        return '전국' if district in ('계', '합계') else None
    if province not in PROVINCES:
        return None
    return province if district == '계' else f'{province}|{district}'


def parse_workbook(path, month):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        total_sheet = next(s for s in workbook if s.title.startswith('시군구별'))
        completed_sheet = next(s for s in workbook if s.title.startswith('공사완료후'))
        rows = list(total_sheet.iter_rows(values_only=True))
        label = norm(rows[2][-1]).strip("'")
        expected = f'{int(month[:4]) % 100}.{int(month[5:])}'
        if label != expected:
            raise ValueError('파일 기준월과 엑셀 마지막 월이 다릅니다')
        total, completed = {}, {}
        for row in rows[3:]:
            key = region_key(row[0], row[1])
            if key and isinstance(row[-1], (int, float)):
                total[key] = int(row[-1])
        for row in list(completed_sheet.iter_rows(values_only=True))[4:]:
            key = region_key(row[0], row[1])
            if key and isinstance(row[2], (int, float)):
                completed[key] = int(row[2])
        for province in ('광주', '전남'):
            keys = districts()[province]
            for data in (total, completed):
                if all(key in data for key in keys):
                    data[province] = sum(data[key] for key in keys)
        if any(p not in total or p not in completed for p in PROVINCES):
            raise ValueError('시도 전체 자료가 없습니다')
        # 시도 총계를 그대로 합산하고 공식 준공 후 전국 값과 교차 확인한다.
        total['전국'] = sum(total[p] for p in PROVINCES)
        if completed.get('전국') != sum(completed[p] for p in PROVINCES):
            raise ValueError('준공 후 전국과 시도 합계가 다릅니다')
        if any(v < 0 or completed.get(k, 0) > v for k, v in total.items() if k in completed):
            raise ValueError('미분양 수치 범위가 맞지 않습니다')
        return {'total': total, 'completed': completed}
    finally:
        workbook.close()


def collect(limit=6):
    with urllib.request.urlopen(SOURCE, timeout=30) as response:
        html = response.read().decode('utf-8')
    matches = re.findall(r"downFile\('([^']+\.xlsx)','([^']+)','([^']+)'", html)
    files = {}
    for label, name, folder in matches:
        date = re.search(r'미분양주택현황\((\d{4})년(\d{1,2})월', label)
        if date:
            month = f'{int(date[1]):04d}-{int(date[2]):02d}'
            files[month] = (label, name, folder)
    if not files:
        raise ValueError('공식 미분양 엑셀을 찾지 못했습니다')
    periods = {}
    for month in sorted(files, reverse=True)[:limit]:
        path = ROOT / 'cache/supply' / (month.replace('-', '') + '.xlsx')
        if not path.exists():
            query = urllib.parse.urlencode(dict(zip(('oFileName', 'rFileName', 'midpath'), files[month])))
            url = 'https://stat.molit.go.kr/portal/common/downLoadFile.do?' + query
            with urllib.request.urlopen(url, timeout=30) as response:
                body = response.read()
            path.parent.mkdir(parents=True, exist_ok=True)
            part = path.with_suffix('.part')
            part.write_bytes(body)
            part.replace(path)
        periods[month] = parse_workbook(path, month)
        print('공식 미분양 추출:', month, flush=True)
    months = sorted(periods)
    keys = sorted({k for period in periods.values() for k in period['total']})
    data = {'schema_version': 1, 'months': months, 'source': SOURCE,
            'collected': dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).strftime('%Y-%m-%d %H:%M KST'),
            'values': {key: {kind: [periods[m][kind].get(key) for m in months] for kind in ('total', 'completed')} for key in keys}}
    dest = ROOT / 'cache/published/housing/supply.json.gz'
    dest.parent.mkdir(parents=True, exist_ok=True)
    part = dest.with_suffix('.part')
    part.write_bytes(gzip.compress(json.dumps(data, ensure_ascii=False, separators=(',', ':')).encode(), mtime=0))
    part.replace(dest)
    return data


if __name__ == '__main__':
    collect()
