"""KB 월간 주택 시계열 엑셀에서 시각화용 데이터를 JSON으로 추출한다.

사용법: python3 scripts/kb_monthly_extract.py "kbdata/202609_월간 주택 시계열.xlsx" out.json ["kbdata/202609_월간 오피스텔 시계열.xlsx"]

월간 파일은 시트마다 머리글 줄 수와 날짜 표기가 다르다(86.1 / '99.1 / 2016.1 / 2009. 1 / 월 숫자만).
날짜는 첫 열을 위에서부터 읽으며 '이전 달 + 1'을 기대값으로 삼아 해석한다.
표 중간의 안내 문구(2022-11 표본 개편 등)는 건너뛰고, 연도가 적힌 날짜가 앞으로 되돌아가면
아래쪽에 다른 표가 시작된 것으로 보고 멈춘다.
"""
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

PROVINCES = ["서울특별시", "부산광역시", "대구광역시", "인천광역시", "(구)광주광역시", "대전광역시", "울산광역시",
             "세종특별자치시", "경기도", "강원특별자치도", "충청북도", "충청남도", "전북특별자치도", "(구)전라남도",
             "경상북도", "경상남도", "제주특별자치도"]
LABEL = {"(구)광주광역시": "광주광역시", "(구)전라남도": "전라남도"}
AGGS = ["전국", "수도권", "6개광역시", "5개광역시", "기타지방"]
STOPS = set(PROVINCES) | set(AGGS) | {"5개광역시(인천外)", "전남광주통합특별시", "제주/서귀포"}
SEOUL_GROUPS = {"강북14개구": "강북 14개 구", "강남11개구": "강남 11개 구"}
REGION_NAMES = set(PROVINCES) | set(AGGS) | {"강북14개구", "강남11개구", "전남광주통합특별시"}
CITY_UNIT = {"경기도", "강원특별자치도", "충청북도", "충청남도", "전북특별자치도", "(구)전라남도", "경상북도", "경상남도"}
TYPES = {"종합": "all", "아파트": "apt", "단독": "detached", "연립": "row"}


def norm(s):
    return re.sub(r"\s+", "", str(s)) if s is not None else ""


def parse_dates(col, restart=False):
    """첫 열 값 목록 → 같은 길이의 'YYYY-MM' 또는 None.

    restart=True면 날짜가 되돌아가는 곳에서 멈추지 않고, 겹치는 앞 구간을 지운 뒤 새 표로 이어 간다
    (HAI처럼 같은 시트에 개편 전·후 표가 위아래로 있는 경우).
    """
    out, prev = [None] * len(col), None
    for i, v in enumerate(col):
        ym = None
        exp = None if prev is None else ((prev[0] + (prev[1] == 12), prev[1] % 12 + 1))
        sv = str(v).strip().strip("'") if v is not None else ""
        if isinstance(v, bool) or v is None:
            ym = None
        elif prev and re.fullmatch(r"\d{1,2}", sv):
            # 월 숫자만 있는 줄은 정확히 다음 달일 때만 인정 (아래쪽 별도 표의 숫자 무시)
            ym = exp if int(sv) == exp[1] else None
        else:
            s = str(v).strip().strip("'").replace("\n", "").replace(" ", "")
            m = re.fullmatch(r"(\d{2}|\d{4})\.(\d{1,2})", s)
            if isinstance(v, float) and not m:
                m = re.fullmatch(r"(\d{2}|\d{4})\.(\d{1,2})", f"{v:.2f}".rstrip("0"))
            if m:
                y = int(m.group(1))
                y = y + (1900 if y >= 50 else 2000) if y < 100 else y
                mo = int(m.group(2))
                # 2016.1(float)은 1월인지 10월인지 알 수 없으므로 순서상 기대값을 우선한다
                if exp and exp[0] == y:
                    mo = exp[1]
                ym = (y, mo) if 1 <= mo <= 12 else None
                if ym and prev and ym <= prev:
                    if not restart:
                        break  # 날짜가 되돌아감 = 아래쪽 별도 표
                    key = f"{ym[0]:04d}-{ym[1]:02d}"
                    out = [None if (x and x >= key) else x for x in out]
        if ym:
            out[i] = f"{ym[0]:04d}-{ym[1]:02d}"
            prev = ym
    return out


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(float(v), 4)
    return None


def rows_of(wb, sheet):
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def table(rows, start, restart=False):
    """(날짜 목록, 데이터 줄 목록)."""
    body = rows[start:]
    dates = parse_dates([r[0] for r in body], restart)
    keep = [(d, r) for d, r in zip(dates, body) if d]
    return [d for d, _ in keep], [r for _, r in keep]


def region_columns(names):
    """열 이름 목록 → {키: 열번호}, 시·도별 하위 지역 목록. 키는 시도 이름 또는 '시도|지역'."""
    keys, views, prov, kids, group, groups = {}, [], None, [], None, {}

    def close():
        if prov and kids:
            views.append({"id": LABEL.get(prov, prov), "unit": "구" if prov == "서울특별시" else ("시" if prov in CITY_UNIT else "구·군"),
                          "children": kids[:], "groups": dict(groups),
                          "groupNames": list(dict.fromkeys(groups.values()))})

    for c, n in enumerate(names):
        if not n:
            continue
        if n in AGGS and n not in keys:
            keys[n] = c
        if n in STOPS:
            close()
            prov, kids, group, groups = (n if n in PROVINCES else None), [], None, {}
            if n in PROVINCES:
                keys[LABEL.get(n, n)] = c
            continue
        if not prov:
            continue
        if n in SEOUL_GROUPS:
            group = SEOUL_GROUPS[n]
            continue
        if prov in CITY_UNIT:
            if n.endswith("구") or n.endswith("군"):
                continue  # 시 아래 구는 생략
            if not n.endswith("시"):
                n += "시"  # 시트에 '의왕', '하남'처럼 '시'가 빠진 이름이 있음
        key = f"{LABEL.get(prov, prov)}|{n}"
        if key in keys:
            continue
        keys[key] = c
        kids.append(key)
        if group:
            groups[key] = group
    close()
    return keys, views


def index_sheet(wb, sheet, name_rows=(1, 2), start=4):
    rows = rows_of(wb, sheet)
    a, b = rows[name_rows[0]], rows[name_rows[-1]]
    names = [norm(b[c]) or norm(a[c]) for c in range(len(a))]
    dates, body = table(rows, start)
    keys, views = region_columns(names)
    vals = {k: [(lambda x: None if x == 0 else x)(num(r[c]) if c < len(r) else None) for r in body] for k, c in keys.items()}
    return dates, vals, views


def region_type_sheet(wb, sheet, start=4, sub_map=None, restart=False):
    """머리글 1줄 = 지역(병합, 앞값 채움), 2줄 = 하위 항목(종합/아파트/... 또는 1분위...)."""
    rows = rows_of(wb, sheet)
    reg, sub = rows[1], rows[2]
    cur, cols = None, {}
    for c in range(1, len(reg)):
        if reg[c] not in (None, ""):
            cur = norm(str(reg[c]).split(" ")[0])
            cur = re.sub(r"[A-Za-z].*$", "", cur)
        s = norm(sub[c])
        if not cur or not s:
            continue
        s = (sub_map or {}).get(s, s)
        if cur not in REGION_NAMES:
            continue  # HAI의 '대출금리', '중위가구' 같은 지역 아닌 열 제외
        cols.setdefault(LABEL.get(cur, cur), {})[s] = c
    dates, body = table(rows, start, restart)
    out = {r: {s: [num(row[c]) if c < len(row) else None for row in body] for s, c in d.items()} for r, d in cols.items()}
    return dates, out


def simple_sheet(wb, sheet, name_row=1, start=3, stop_word=None):
    rows = rows_of(wb, sheet)
    names = [norm(x) for x in rows[name_row]]
    dates, body = table(rows, start)
    out = {}
    for c, n in enumerate(names):
        if c == 0 or not n or n in ("구분", "지역", "년도"):
            continue
        out[LABEL.get(n, n)] = [num(r[c]) if c < len(r) else None for r in body]
    return dates, out


def series_obj(dates, vals):
    return {"dates": dates, "values": vals}


OT_REGION = {"종합지역": "전국", "5개광역시주)": "5개광역시", "서울": "서울특별시", "인천": "인천광역시", "경기": "경기도"}
OT_SIZE_GROUPS = ["전국", "수도권", "서울특별시", "경기도"]  # 면적별 블록 순서 (통계개요의 공표범위)


def officetel(path):
    """KB 월간 오피스텔 시계열: 시트마다 같은 열 구성(권역 11개 + 면적 4블록×5구간)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {"source": path.split("/")[-1]}
    for sheet, key in [("매매가격지수", "index"), ("매매평균가격", "mean_sale"), ("전세평균가격", "mean_jeonse"),
                       ("매매전세비율", "ratio"), ("임대수익률", "yield")]:
        rows = rows_of(wb, sheet)
        a, b = rows[2], rows[3]
        dates, body = table(rows, 4)
        regions, sizes, block = {}, {}, -1
        for c in range(1, len(a)):
            n = norm(b[c]) or norm(a[c])
            if not n:
                continue
            col = [num(r[c]) if c < len(r) else None for r in body]
            if n == "초소형":
                block += 1
            if block >= 0:
                sizes.setdefault(OT_SIZE_GROUPS[block], {})[n] = col
            else:
                regions[OT_REGION.get(n, n)] = col
        out[key] = {"dates": dates, "values": regions, "sizes": sizes}
    out["asof"] = out["index"]["dates"][-1]
    return out


def check(data):
    """추출 결과 점검. KB가 시트 구조를 바꿨을 때 엉뚱한 화면이 게시되지 않도록 빌드를 멈춘다."""
    problems = []

    def contiguous(name, ds, since=None):
        n = [int(x[:4]) * 12 + int(x[5:]) for x in ds if not since or x >= since]
        if any(b - a != 1 for a, b in zip(n, n[1:])):
            problems.append(f"{name}: 월이 끊기거나 겹침")

    if len(data["index"]["sale_apt"]["values"]) < 140:
        problems.append(f"아파트 매매지수 지역 수가 {len(data['index']['sale_apt']['values'])}개뿐 (140개 이상이어야 함)")
    if len(data["views"]) < 14:
        problems.append(f"시·도별 하위 지역 묶음이 {len(data['views'])}개뿐")
    for k in ["jeonse_ratio", "mean_sale", "median_sale", "outlook_sale", "leading50"]:
        contiguous(k, data[k]["dates"])
        if data[k]["dates"][-1] != data["asof"]:
            problems.append(f"{k}: 마지막 달 {data[k]['dates'][-1]}이 기준월 {data['asof']}과 다름")
    contiguous("index", data["index"]["sale_apt"]["dates"], since="2000-01")
    if "officetel" in data:
        contiguous("officetel", data["officetel"]["index"]["dates"])
        if len(data["officetel"]["index"]["values"]) < 8:
            problems.append("오피스텔 권역 수가 너무 적음")
    if problems:
        raise SystemExit("월간 자료 점검 실패:\n - " + "\n - ".join(problems))


def main(path, out, officetel_path=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {"source": path.split("/")[-1]}

    # 1~8. 유형별 매매·전세 지수 (2026.01 = 100)
    idx, views = {}, None
    for sheet, key in [("1.매매종합", "sale_all"), ("2.매매APT", "sale_apt"), ("3.매매단독", "sale_detached"), ("4.매매연립", "sale_row"),
                       ("5.전세종합", "jeonse_all"), ("6.전세APT", "jeonse_apt"), ("7.전세단독", "jeonse_detached"), ("8.전세연립", "jeonse_row")]:
        d, v, vw = index_sheet(wb, sheet)
        idx[key] = series_obj(d, v)
        if sheet == "2.매매APT":
            views = vw
    data["index"], data["views"] = idx, views
    data["asof"] = idx["sale_apt"]["dates"][-1]

    # 28. 아파트 매매가 대비 전세가 비율(%)
    d, v, _ = index_sheet(wb, "28.아파트매매전세비", name_rows=(1,), start=3)
    data["jeonse_ratio"] = series_obj(d, v)

    # 41·42·43·44. 평균·중위 가격(만원) — 지역 × 유형
    for sheet, key in [("41.평균매매", "mean_sale"), ("42.평균전세", "mean_jeonse"), ("43.중위매매", "median_sale"), ("44.중위전세", "median_jeonse")]:
        d, v = region_type_sheet(wb, sheet, sub_map=TYPES)
        data[key] = series_obj(d, v)

    # 14. 주택구매력지수(HAI) — 지역 × 유형
    d, v = region_type_sheet(wb, "14.NEW_HAI", sub_map=TYPES, restart=True)  # 2019년부터 개편 소득 기준 표
    data["hai"] = series_obj(d, v)

    # 25·26. 3개월 뒤 가격 전망지수 — 지역별 '전망지수' 열만
    for sheet, key in [("25.KB부동산 매매가격 전망지수", "outlook_sale"), ("26.KB부동산 전세가격 전망지수", "outlook_jeonse")]:
        d, v = region_type_sheet(wb, sheet)
        data[key] = series_obj(d, {r: next((s for n, s in sub.items() if "전망지수" in n), None) for r, sub in v.items()})

    # 53·54. 5분위 평균 아파트 가격(만원)과 5분위 배율
    for sheet, key in [("53.5분위(아파트매매)", "quintile_sale"), ("54.5분위(아파트전세)", "quintile_jeonse")]:
        d, v = region_type_sheet(wb, sheet)
        data[key] = series_obj(d, v)

    # 9. 월세지수, 59. 전월세전환율 (수도권 위주)
    data["rent"] = series_obj(*simple_sheet(wb, "9.KB아파트 월세지수", start=4))
    data["conversion"] = series_obj(*simple_sheet(wb, "59.전월세전환율", start=3))

    # 16. KB 선도아파트 50
    rows = rows_of(wb, "16.선도50")
    d, body = table(rows, 3)
    data["leading50"] = series_obj(d, {"지수": [num(r[1]) for r in body]})

    if officetel_path:
        data["officetel"] = officetel(officetel_path)
    check(data)

    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    if officetel_path:
        print(f"officetel       {data['officetel']['index']['dates'][0]} ~ {data['officetel']['asof']} 권역 {list(data['officetel']['index']['values'])} 면적 {list(data['officetel']['index']['sizes'])}")
    for k, v in data.items():
        if isinstance(v, dict) and "dates" in v and v["dates"]:
            print(f"{k:15s} {v['dates'][0]} ~ {v['dates'][-1]} ({len(v['dates'])}개월) 항목 {len(v['values'])}")
        elif k == "index":
            for kk, vv in v.items():
                print(f"index.{kk:13s} {vv['dates'][0]} ~ {vv['dates'][-1]} ({len(vv['dates'])}개월) 지역 {len(vv['values'])}")
    print("views:", [(x["id"], len(x["children"])) for x in data["views"]])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
